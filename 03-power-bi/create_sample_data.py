"""Script to generate the sample dashboard_lanas_excel.xlsx data file."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path


def create_dashboard_lanas_excel(output_path: str | Path = "dashboard_lanas_excel.xlsx") -> Path:
    """
    Create a sample Excel workbook with lanas (yarn) business data
    that serves as the data source for the Power BI dashboard.

    Args:
        output_path: Path where the Excel file will be saved.

    Returns:
        Path to the created file.
    """
    output_path = Path(output_path)
    wb = openpyxl.Workbook()

    # --- Sheet 1: Ventas (Sales) ---
    ws_ventas = wb.active
    ws_ventas.title = "Ventas"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center = Alignment(horizontal="center")

    ventas_headers = ["Fecha", "Producto", "Color", "Gramaje_g", "Cantidad", "Precio_Unitario", "Total"]
    for col, header in enumerate(ventas_headers, start=1):
        cell = ws_ventas.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    ventas_data = [
        ["2024-01-05", "Merino Fino",    "Azul Marino",   100, 10, 850.0,  8500.0],
        ["2024-01-12", "Alpaca Premium", "Natural",       200, 5,  1200.0, 6000.0],
        ["2024-01-18", "Cotton Mix",     "Blanco",        150, 20, 450.0,  9000.0],
        ["2024-02-03", "Merino Fino",    "Rojo Vino",     100, 8,  850.0,  6800.0],
        ["2024-02-14", "Mohair Luxe",    "Rosa Pastel",   50,  3,  1800.0, 5400.0],
        ["2024-02-20", "Cotton Mix",     "Verde Salvia",  150, 15, 450.0,  6750.0],
        ["2024-03-07", "Alpaca Premium", "Gris Perla",    200, 7,  1200.0, 8400.0],
        ["2024-03-15", "Merino Fino",    "Amarillo",      100, 12, 850.0,  10200.0],
        ["2024-03-22", "Mohair Luxe",    "Lavanda",       50,  4,  1800.0, 7200.0],
        ["2024-04-01", "Cotton Mix",     "Naranja",       150, 18, 450.0,  8100.0],
        ["2024-04-10", "Merino Fino",    "Negro",         100, 6,  850.0,  5100.0],
        ["2024-04-25", "Alpaca Premium", "Beige",         200, 9,  1200.0, 10800.0],
    ]
    for row_data in ventas_data:
        ws_ventas.append(row_data)

    for col in ws_ventas.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_ventas.column_dimensions[col[0].column_letter].width = max_len + 4

    # --- Sheet 2: Inventario (Inventory) ---
    ws_inv = wb.create_sheet("Inventario")

    inv_headers = ["Producto", "Color", "Gramaje_g", "Stock_Unidades", "Costo_Unitario", "Valor_Total"]
    for col, header in enumerate(inv_headers, start=1):
        cell = ws_inv.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    inventario_data = [
        ["Merino Fino",    "Azul Marino",  100, 45,  420.0, 18900.0],
        ["Merino Fino",    "Rojo Vino",    100, 30,  420.0, 12600.0],
        ["Merino Fino",    "Amarillo",     100, 28,  420.0, 11760.0],
        ["Merino Fino",    "Negro",        100, 50,  420.0, 21000.0],
        ["Alpaca Premium", "Natural",      200, 20,  600.0, 12000.0],
        ["Alpaca Premium", "Gris Perla",   200, 15,  600.0, 9000.0],
        ["Alpaca Premium", "Beige",        200, 25,  600.0, 15000.0],
        ["Cotton Mix",     "Blanco",       150, 80,  220.0, 17600.0],
        ["Cotton Mix",     "Verde Salvia", 150, 60,  220.0, 13200.0],
        ["Cotton Mix",     "Naranja",      150, 40,  220.0, 8800.0],
        ["Mohair Luxe",    "Rosa Pastel",  50,  12,  900.0, 10800.0],
        ["Mohair Luxe",    "Lavanda",      50,  8,   900.0, 7200.0],
    ]
    for row_data in inventario_data:
        ws_inv.append(row_data)

    for col in ws_inv.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_inv.column_dimensions[col[0].column_letter].width = max_len + 4

    # --- Sheet 3: Proveedores (Suppliers) ---
    ws_prov = wb.create_sheet("Proveedores")

    prov_headers = ["Proveedor", "Pais", "Producto", "Plazo_Entrega_Dias", "Calificacion"]
    for col, header in enumerate(prov_headers, start=1):
        cell = ws_prov.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    proveedores_data = [
        ["LanasPaul S.A.",     "Argentina", "Merino Fino",    15, 4.8],
        ["AlpacaAndes Ltda.",  "Peru",      "Alpaca Premium", 30, 4.9],
        ["CottonCo.",          "Brasil",    "Cotton Mix",     10, 4.5],
        ["MohairEurope GmbH.", "Alemania",  "Mohair Luxe",    45, 4.7],
    ]
    for row_data in proveedores_data:
        ws_prov.append(row_data)

    for col in ws_prov.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_prov.column_dimensions[col[0].column_letter].width = max_len + 4

    wb.save(output_path)
    print(f"Created: {output_path}")
    return output_path


if __name__ == "__main__":
    script_dir = Path(__file__).parent
    create_dashboard_lanas_excel(script_dir / "dashboard_lanas_excel.xlsx")
